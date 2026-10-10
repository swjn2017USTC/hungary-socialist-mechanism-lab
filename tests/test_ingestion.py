"""Offline ingestion tests: synthetic payloads, no network.

Every fixture mimics the audited real structure (quoted multiline KSH CSV
headers, div#thead/#tbody HTML tables, XLSX via openpyxl, zipped SDMX XML)
so the parsers, cache, rights gate and artifact registration are exercised
end-to-end without a single network call.
"""

from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

import pytest
import yaml

from hsm_lab.contract.schemas import ConfigError, ResearchGateError, TargetUse
from hsm_lab.data.adapters import (
    SourceContentIncomplete,
    SourceSchemaChanged,
    parse_fed_h15_zip,
    parse_fred_fedfunds_csv,
    parse_ksh_agriculture_html,
    parse_ksh_gdp_csv,
    parse_ksh_industry_csv,
    parse_ksh_trade_html,
    parse_mnb_fx_html,
    parse_worldbank_pink_sheet_xlsx,
)
from hsm_lab.data.ingestion import (
    IngestionRequest,
    _request_hash,
    ingest_source,
)
from tests.conftest import REPO_ROOT

# ----------------------------------------------------------------------
# Fixture payloads (audited real structure, synthetic values)
# ----------------------------------------------------------------------

GDP_CSV = (
    "21.1.1.2. Long time series of GDP, consumption and capital formation;;;;;\r\n"
    'Year;"Gross domestic product (GDP), \n1960 = 100%";'
    '"Actual final consumption total, \n1960 = 100%";'
    '"Actual final consumption, of which actual final consumption of households, \n1960 = 100%";'
    '"Gross capital formation total, \n1960 = 100%";'
    '"Gross capital formation, of which fixed capital formation,  \n1960 = 100%"\r\n'
    + "\n".join(
        f"{y};{100 + y - 1960};{100 + 2 * (y - 1960)};{100 + 3 * (y - 1960)};"
        f"{100 + 4 * (y - 1960)};{100 + 5 * (y - 1960)}"
        for y in range(1960, 1986)
    )
    + "\r\n"
)

INDUSTRY_HEADER = (
    'Year;"Volume indices of production,\n1960=100.0%";"Electricity, \nmillion kWh";'
    '"Coal,\nthousand tonnes";"Natural gas,\nmillion m3";Bauxite thousand tonnes;'
    "Crude steel thousand tonnes;Cement thousand tonnes;\"Buses,\npieces\";"
    '"Television sets,\nthousands";Sulphuric acid, technical thousand tonnes;'
    "Mineral or chemical fertilizers thousand tonnes;Plastics in primary form thousand tonnes;"
    "Paper and card board thousand tonnes;\"Cotton fabrics, \nmillion m2\";"
    '"Foot-wear,\nthousand pairs";"Sugar,\nthousand tonnes";"Beer,\nmillion litres"'
)


def industry_csv(years=range(1960, 1986), *, header=INDUSTRY_HEADER, extra_col=False):
    title = (
        "13.1.1.2. Long time series of the volume index of industrial production "
        "and of production of some significant industrial products ;;;;;;;;;;;;;;;;\r\n"
    )
    rows = []
    for y in years:
        base = (
            f"{y};{100 + (y - 1960) * 9}"
            ";7 617;26 524;342;1 190;1 887;1 571;1 877;139;164;102;10;138;247;21 334;380;356"
        )
        if extra_col:
            base += ";999"
        rows.append(base)
    return title + header + "\r\n" + "\r\n".join(rows) + "\r\n"


AGRI_HTML = (
"""<!DOCTYPE html>
<html><head><meta http-equiv="Content-Type" content="text/html; charset=iso-8859-2" />
<title>STADAT</title></head><body>
<div id="title"><h1>4.1.1. Major data of agriculture (1960&ndash;)"""
"""<span class="pagenum">(2/3)</span></h1></div>
<div id="table">
\t<div id="thead">
\t\t<table style="height: 68px;">
\t\t\t<tr>
\t\t\t\t<th rowspan="3" width="10%">Year</th>
\t\t\t\t<th rowspan="2" width="10%">Cereals</th>
\t\t\t\t<th colspan="2" width="20%">Of which</th>
\t\t\t\t<th rowspan="2" width="10%">Sugarbeet</th>
\t\t\t\t<th rowspan="2" width="10%">Oilseeds</th>
\t\t\t\t<th rowspan="2" width="10%">Potatoes</th>
\t\t\t\t<th rowspan="2" width="10%">Vegetables<span class="idx" title="paprika note">c</span></th>
\t\t\t\t<th rowspan="2" width="10%">Fruits<span class="idx" title="strawberry">d</span></th>
\t\t\t\t<th rowspan="2" width="10%">Grapes</th>
\t\t\t</tr>
\t\t\t<tr><th width="10%">wheat</th><th width="10%">maize</th></tr>
\t\t\t<tr><th colspan="9" width="90%">thousand tons</th></tr>
\t\t</table>
\t</div>
\t<div id="tbody">
\t\t<table>
\t\t\t<tr><td class="oc">1960</td><td>6 960</td><td>1 803</td><td>3 534</td>"""
"""<td>3 370</td><td>100</td><td>3 001</td><td>1 374</td>"""
"""<td>737</td><td>491</td></tr>
\t\t\t<tr><td class="oc">1961</td><td>6 210</td><td>1 974</td><td>2 737</td>"""
"""<td>2 356</td><td>133</td><td>1 830</td><td>1 254</td>"""
"""<td>1 032</td><td>591</td></tr>
\t\t</table>
\t</div>
</div>
</body></html>
"""
)


def agri_html(years=(1960, 1961, 1962)):
    head = AGRI_HTML.split("\t<div id=\"tbody\">")[0] + '\t<div id="tbody">\n\t\t<table>\n'
    body = []
    for y in years:
        body.append(
            f"\t\t\t<tr><td class=\"oc\">{y}</td><td>6 960</td><td>1 803</td>"
            "<td>3 534</td><td>3 370</td><td>100</td><td>3 001</td>"
            "<td>1 374</td><td>737</td><td>491</td></tr>"
        )
    return head + "\n".join(body) + "\n\t\t</table>\n\t</div>\n</div>\n</body></html>\n"


TRADE_HTML = (
"""<html><body>
<div id="title"><h1>3.3. External trade in goods (1960&ndash;)</h1></div>
<div id="table">
\t<div id="thead"><table>
\t\t<tr><th rowspan="3" width="10%">Year</th>"""
"""<th colspan="2" width="90%">Volume indices of external trade in goods</th></tr>
\t\t<tr><th width="45%">imports</th><th width="45%">exports</th></tr>
\t\t<tr><th colspan="2" width="90%">1960 = 100%</th></tr>
\t</table></div>
\t<div id="tbody"><table>
\t\t<tr><td class="oc">1960</td><td>100</td><td>100</td></tr>
\t\t<tr><td class="oc">1961</td><td>108</td><td>118</td></tr>
\t</table></div>
</div>
</body></html>
"""
)


def mnb_html(quote_rows=("03 January 1985", "51.30", "30 March 1985", "50.88")):
    rows = "\n".join(
        f'\t\t\t<tr><td class="rotate">{quote_rows[i]}</td>'
        f'<td class="rotate">{quote_rows[i + 1]}</td></tr>'
        for i in range(0, len(quote_rows), 2)
    )
    return f"""<html><body><table class="pricetable">
\t\t\t<thead>
\t\t\t\t<tr><th></th><th class="rotate">USD</th></tr>
\t\t\t\t<tr><th></th><th class="rotate">US Dollar</th></tr>
\t\t\t\t<tr><th></th><th class="rotate">1</th></tr>
\t\t\t</thead>
\t\t\t<tbody>
{rows}
\t\t\t</tbody>
\t\t</table></body></html>
"""


FRED_CSV = (
    "observation_date,FEDFUNDS\n"
    + "\n".join(f"{y}-{m:02d}-01,2.5" for y in (1957,) for m in (1, 2, 3))
    + "\n"
)


def fed_xml(months=(("1957-01-31", "2.84"), ("1957-02-28", "3.00"), ("1957-03-31", "2.96"))):
    obs = "\n".join(
        f'      <frb:Obs OBS_STATUS="A" OBS_VALUE="{v}" TIME_PERIOD="{p}" />'
        for p, v in months
    )
    return (
        """<?xml version="1.0" encoding="UTF-8" standalone="no"?>
<message:MessageGroup xmlns:message="http://www.SDMX.org/resources/SDMXML/schemas/v1_0/message" """
"""xmlns:frb="http://www.federalreserve.gov/structure/compact/common" """
"""xmlns:kf="http://www.federalreserve.gov/structure/compact/keyfamily">
  <kf:Series CURRENCY="NA" FREQ="9" INSTRUMENT="FF" MATURITY="O" """
"""SERIES_NAME="RIFSPFF_N.B" UNIT="Percent:_Per_Year" UNIT_MULT="1">
    <frb:Obs OBS_STATUS="A" OBS_VALUE="9.99" TIME_PERIOD="1957-01-31" />
  </kf:Series>
  <kf:Series CURRENCY="NA" FREQ="129" INSTRUMENT="FF" MATURITY="O" """
"""SERIES_NAME="RIFSPFF_N.M" UNIT="Percent:_Per_Year" UNIT_MULT="1">
"""
        + obs
        + """
  </kf:Series>
</message:MessageGroup>
"""
    )


def fed_zip_bytes(xml: str) -> bytes:
    import io

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("H15_data.xml", xml)
    return buffer.getvalue()


def write_wb_xlsx(
    path: Path,
    values=(("1960M01", 1.63000011444), ("1960M02", 1.63000011444), ("1960M03", 1.65)),
):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Monthly Prices"
    sheet["A1"] = "World Bank Commodity Price Data (The Pink Sheet)"
    sheet["A5"] = None
    sheet["B5"] = "Crude oil, average"
    sheet["C5"] = "Crude oil, Brent"
    sheet["B6"] = "($/bbl)"
    sheet["C6"] = "($/bbl)"
    for i, (period, value) in enumerate(values, start=7):
        sheet.cell(row=i, column=1, value=period)
        sheet.cell(row=i, column=2, value=value)
    workbook.save(path)


# ----------------------------------------------------------------------
# Parser acceptance (real structure, plan acceptance values)
# ----------------------------------------------------------------------


def test_ksh_gdp_csv_accepts_multiline_headers(tmp_path):
    payload = tmp_path / "gdp0002.csv"
    payload.write_text(GDP_CSV, encoding="utf-8")
    series = {s.variable_id: s for s in parse_ksh_gdp_csv(payload, start_year=1960, end_year=1985)}
    assert set(series) == {"gdp", "consumption", "investment"}
    assert series["gdp"].unit == "index_1960_100"
    assert dict(series["gdp"].rows)["1960"] == 100.0
    assert dict(series["gdp"].rows)["1985"] == 125.0
    assert dict(series["consumption"].rows)["1961"] == 103.0


def test_ksh_industry_csv_acceptance_values(tmp_path):
    payload = tmp_path / "ipa0002.csv"
    payload.write_text(industry_csv(), encoding="utf-8")
    series = {
        s.variable_id: s
        for s in parse_ksh_industry_csv(payload, start_year=1960, end_year=1985)
    }
    assert set(series) == {
        "industrial_output",
        "industrial_electricity",
        "industrial_crude_steel",
        "industrial_fertilizer",
        "industrial_cement",
        "industrial_television",
    }
    values = {k: dict(v.rows)["1960"] for k, v in series.items()}
    assert values["industrial_electricity"] == 7617.0
    assert values["industrial_crude_steel"] == 1887.0
    assert values["industrial_fertilizer"] == 102.0
    assert values["industrial_cement"] == 1571.0
    assert values["industrial_television"] == 139.0
    assert values["industrial_output"] == 100.0
    assert series["industrial_electricity"].unit == "million_kWh"
    assert series["industrial_television"].unit == "thousand_units"


def test_ksh_agriculture_html_acceptance_values(tmp_path):
    payload = tmp_path / "h_omf001b.html"
    payload.write_text(agri_html(range(1960, 1963)), encoding="utf-8")
    series = {
        s.variable_id: s
        for s in parse_ksh_agriculture_html(payload, start_year=1960, end_year=1962)
    }
    assert dict(series["ag_cereals"].rows)["1960"] == 6960.0
    assert dict(series["ag_wheat"].rows)["1960"] == 1803.0
    assert dict(series["ag_maize"].rows)["1960"] == 3534.0
    assert all(s.unit == "thousand_tonnes" for s in series.values())


def test_ksh_trade_html_acceptance_values(tmp_path):
    payload = tmp_path / "h_qkt001.html"
    payload.write_text(TRADE_HTML, encoding="utf-8")
    series = {
        s.variable_id: s
        for s in parse_ksh_trade_html(payload, start_year=1960, end_year=1961)
    }
    assert dict(series["trade_import_volume_index"].rows)["1961"] == 108.0
    assert dict(series["trade_export_volume_index"].rows)["1961"] == 118.0
    assert series["trade_import_volume_index"].unit == "index_1960_100"


def test_worldbank_xlsx_acceptance_values(tmp_path):
    payload = tmp_path / "CMO.xlsx"
    write_wb_xlsx(payload)
    series = parse_worldbank_pink_sheet_xlsx(payload, start_month="1960-01", end_month="1960-03")
    assert series[0].variable_id == "world_oil_price"
    assert series[0].unit == "USD_per_barrel"
    assert dict(series[0].rows)["1960-01"] == 1.63000011444


def test_fed_zip_acceptance_values(tmp_path):
    payload = tmp_path / "FRB_h15_xml.zip"
    payload.write_bytes(fed_zip_bytes(fed_xml()))
    series = parse_fed_h15_zip(payload, start_month="1957-01", end_month="1957-03")
    rows = dict(series[0].rows)
    assert rows == {"1957-01": 2.84, "1957-02": 3.0, "1957-03": 2.96}
    mean_q1 = sum(rows.values()) / 3
    assert abs(mean_q1 - 2.933333333333333) < 1e-12
    assert series[0].unit == "percent_per_year"


def test_mnb_html_never_divides_by_100(tmp_path):
    payload = tmp_path / "mnb.html"
    payload.write_text(mnb_html(), encoding="utf-8")
    series = parse_mnb_fx_html(payload, start_date="1985-01-01", end_date="1985-12-31")
    rows = dict(series[0].rows)
    assert rows["1985-01-03"] == 51.30
    assert rows["1985-03-30"] == 50.88  # page is already HUF per 1 USD
    assert series[0].unit == "HUF_per_USD"


def test_fred_gated_parser_fixture(tmp_path):
    payload = tmp_path / "fredgraph.csv"
    payload.write_text(FRED_CSV, encoding="utf-8")
    series = parse_fred_fedfunds_csv(payload, start_month="1957-01", end_month="1957-03")
    assert dict(series[0].rows) == {"1957-01": 2.5, "1957-02": 2.5, "1957-03": 2.5}


# ----------------------------------------------------------------------
# Schema drift and fail-closed content
# ----------------------------------------------------------------------


def test_extra_column_is_drift(tmp_path):
    payload = tmp_path / "ipa0002.csv"
    payload.write_text(industry_csv(extra_col=True), encoding="utf-8")
    with pytest.raises(SourceSchemaChanged, match="source_schema_changed.*ksh_industry"):
        parse_ksh_industry_csv(payload, start_year=1960, end_year=1985)


def test_renamed_header_is_drift(tmp_path):
    header = INDUSTRY_HEADER.replace("Crude steel thousand tonnes", "Pig iron thousand tonnes")
    payload = tmp_path / "ipa0002.csv"
    payload.write_text(industry_csv(header=header), encoding="utf-8")
    with pytest.raises(SourceSchemaChanged, match="drifted"):
        parse_ksh_industry_csv(payload, start_year=1960, end_year=1985)


def test_html_error_page_is_drift_not_csv(tmp_path):
    payload = tmp_path / "gdp0002.csv"
    payload.write_text("<html><body>Error 500</body></html>", encoding="utf-8")
    with pytest.raises(SourceSchemaChanged, match="no header row"):
        parse_ksh_gdp_csv(payload, start_year=1960, end_year=1985)


def test_missing_year_in_core_series_fails_closed(tmp_path):
    payload = tmp_path / "ipa0002.csv"
    payload.write_text(
        industry_csv(years=[y for y in range(1960, 1986) if y != 1971]), encoding="utf-8"
    )
    with pytest.raises(SourceContentIncomplete, match="missing year"):
        parse_ksh_industry_csv(payload, start_year=1960, end_year=1985)


def test_missing_marker_in_core_series_fails_closed(tmp_path):
    text = industry_csv().replace("1971;199;7 617", "1971;199;...")
    payload = tmp_path / "ipa0002.csv"
    payload.write_text(text, encoding="utf-8")
    with pytest.raises(SourceContentIncomplete, match="missing year"):
        parse_ksh_industry_csv(payload, start_year=1960, end_year=1985)


def test_wb_wrong_unit_row_is_drift(tmp_path):
    payload = tmp_path / "CMO.xlsx"
    write_wb_xlsx(payload)
    import openpyxl

    workbook = openpyxl.load_workbook(payload)
    workbook["Monthly Prices"]["B6"] = "($/mt)"
    workbook.save(payload)
    with pytest.raises(SourceSchemaChanged, match="unit row"):
        parse_worldbank_pink_sheet_xlsx(payload, start_month="1960-01", end_month="1960-03")


def test_wb_missing_month_fails_closed(tmp_path):
    payload = tmp_path / "CMO.xlsx"
    write_wb_xlsx(payload, values=(("1960M01", 1.63), ("1960M03", 1.65)))
    with pytest.raises(SourceContentIncomplete, match="missing month"):
        parse_worldbank_pink_sheet_xlsx(payload, start_month="1960-01", end_month="1960-03")


def test_wb_duplicate_crude_header_is_drift(tmp_path):
    payload = tmp_path / "CMO.xlsx"
    write_wb_xlsx(payload)
    import openpyxl

    workbook = openpyxl.load_workbook(payload)
    workbook["Monthly Prices"]["D5"] = "Crude oil, average"
    workbook.save(payload)
    with pytest.raises(SourceSchemaChanged, match="not unique"):
        parse_worldbank_pink_sheet_xlsx(payload, start_month="1960-01", end_month="1960-03")


def test_fed_missing_month_fails_closed(tmp_path):
    payload = tmp_path / "FRB_h15_xml.zip"
    payload.write_bytes(
        fed_zip_bytes(fed_xml(months=(("1957-01-31", "2.84"), ("1957-03-31", "2.96"))))
    )
    with pytest.raises(SourceContentIncomplete, match="missing month"):
        parse_fed_h15_zip(payload, start_month="1957-01", end_month="1957-03")


def test_fed_doctype_refused_before_parsing(tmp_path):
    xml = (
        "<!DOCTYPE message:MessageGroup [<!ENTITY xxe SYSTEM 'file:///etc/passwd'>]>\n"
        + fed_xml()
    )
    payload = tmp_path / "FRB_h15_xml.zip"
    payload.write_bytes(fed_zip_bytes(xml))
    with pytest.raises(SourceSchemaChanged, match="DTD"):
        parse_fed_h15_zip(payload, start_month="1957-01", end_month="1957-03")


def test_fed_wrong_series_attributes_refused(tmp_path):
    xml = fed_xml().replace('FREQ="129"', 'FREQ="130"')
    payload = tmp_path / "FRB_h15_xml.zip"
    payload.write_bytes(fed_zip_bytes(xml))
    with pytest.raises(SourceSchemaChanged, match="exactly one Series"):
        parse_fed_h15_zip(payload, start_month="1957-01", end_month="1957-03")


def test_mnb_wrong_thead_is_drift(tmp_path):
    payload = tmp_path / "mnb.html"
    payload.write_text(mnb_html().replace(">USD<", ">EUR<"), encoding="utf-8")
    with pytest.raises(SourceSchemaChanged, match="USD / US Dollar / 1"):
        parse_mnb_fx_html(payload, start_date="1985-01-01", end_date="1985-12-31")


def test_mnb_empty_result_fails_closed(tmp_path):
    payload = tmp_path / "mnb.html"
    payload.write_text(mnb_html(quote_rows=()), encoding="utf-8")
    with pytest.raises(SourceContentIncomplete, match="no USD quotes"):
        parse_mnb_fx_html(payload, start_date="1985-01-01", end_date="1985-12-31")


def test_agri_wrong_unit_row_is_drift(tmp_path):
    payload = tmp_path / "h_omf001b.html"
    payload.write_text(
        agri_html(range(1960, 1963)).replace("thousand tons", "million tons"),
        encoding="utf-8",
    )
    with pytest.raises(SourceSchemaChanged, match="unit row"):
        parse_ksh_agriculture_html(payload, start_year=1960, end_year=1962)


def test_fred_bad_header_is_drift(tmp_path):
    payload = tmp_path / "fredgraph.csv"
    payload.write_text("date,value\n1960-01-01,1\n", encoding="utf-8")
    with pytest.raises(SourceSchemaChanged, match="header"):
        parse_fred_fedfunds_csv(payload, start_month="1960-01", end_month="1960-01")


# ----------------------------------------------------------------------
# Orchestration: rights gate, cache, offline, refresh (monkeypatched net)
# ----------------------------------------------------------------------


def _make_repo(tmp_path: Path) -> Path:
    root = tmp_path / "repo"
    (root / "config").mkdir(parents=True)
    (root / "data" / "manifests").mkdir(parents=True)
    shutil.copy(REPO_ROOT / "config" / "research_contract.json", root / "config")
    shutil.copy(REPO_ROOT / "data" / "manifests" / "sources.yml", root / "data" / "manifests")
    shutil.copy(REPO_ROOT / "data" / "manifests" / "variables.yml", root / "data" / "manifests")
    # The live repository may contain admitted artifacts whose normalized
    # files are intentionally ignored; isolated downloader tests start empty.
    (root / "data" / "manifests" / "artifacts.yml").write_text(
        "schema_version: '1'\nartifacts: []\n", encoding="utf-8"
    )
    return root


class _FakeResponse:
    def __init__(self, data: bytes, url: str):
        self._data = data
        self.status = 200
        self._url = url
        self.headers = {"Content-Type": "text/plain", "Content-Length": str(len(data))}

    def read(self, size=-1):
        if not self._data:
            return b""
        chunk, self._data = self._data[: max(size, 1)], self._data[max(size, 1) :]
        return chunk

    def geturl(self):
        return self._url

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False


def _install_net(monkeypatch, responses: dict[str, bytes], calls: list[str]):
    import urllib.request

    def fake_urlopen(request, timeout=None):
        calls.append(request.full_url)
        if request.full_url not in responses:
            raise AssertionError(f"unexpected network call to {request.full_url}")
        return _FakeResponse(responses[request.full_url], request.full_url)

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)


KSH_GDP_URL = "https://www.ksh.hu/stadat_files/gdp/en/gdp0002.csv"


def test_offline_cache_miss_raises_without_network(tmp_path, monkeypatch):
    root = _make_repo(tmp_path)
    calls: list[str] = []
    _install_net(monkeypatch, {}, calls)
    with pytest.raises(ResearchGateError, match="offline cache miss"):
        ingest_source(root, IngestionRequest("ksh_gdp", "1960", "1985", offline=True))
    assert calls == []


def test_offline_and_refresh_is_config_error(tmp_path):
    root = _make_repo(tmp_path)
    with pytest.raises(ConfigError, match="mutually exclusive"):
        ingest_source(root, IngestionRequest("ksh_gdp", "1960", "1985", offline=True, refresh=True))


def test_unknown_source_is_config_error(tmp_path):
    root = _make_repo(tmp_path)
    with pytest.raises(ConfigError, match="unknown ingestion source"):
        ingest_source(root, IngestionRequest("not_a_source", "1960", "1985"))


def test_fred_rights_gate_blocks_before_network_and_cache(tmp_path, monkeypatch):
    root = _make_repo(tmp_path)
    # Pre-seed a poisoned cache: even a present cache entry must not be read.
    payload = b"observation_date,FEDFUNDS\n"
    import hashlib

    digest = hashlib.sha256(payload).hexdigest()
    content = root / "data" / "raw" / "fred_fedfunds" / digest
    content.mkdir(parents=True)
    (content / "payload.csv").write_bytes(payload)
    (content / "retrieval.json").write_text(
        json.dumps(
            {
                "schema_version": "1",
                "source_id": "fred_fedfunds",
                "retrieved_at_utc": "2026-10-07T00:00:00+00:00",
                "response": {},
                "raw_sha256": digest,
            }
        ),
        encoding="utf-8",
    )
    (root / "data" / "raw").mkdir(parents=True, exist_ok=True)
    (root / "data" / "raw" / "cache-index.json").write_text(
        json.dumps(
            {
                "schema_version": "1",
                "entries": {
                    _request_hash(
                        "https://fred.stlouisfed.org/graph/fredgraph.csv"
                        "?id=FEDFUNDS&cosd=1957-01-01&coed=1985-12-31"
                    ): digest
                },
            }
        ),
        encoding="utf-8",
    )
    calls: list[str] = []
    _install_net(monkeypatch, {}, calls)
    with pytest.raises(ResearchGateError, match="rights-blocked"):
        ingest_source(root, IngestionRequest("fred_fedfunds", "1957-01", "1985-12"))
    assert calls == []


def _manifest_artifact_ids(root: Path) -> set[str]:
    manifest = yaml.safe_load((root / "data" / "manifests" / "artifacts.yml").read_text())
    return {entry["artifact_id"] for entry in manifest["artifacts"]}


def test_download_cache_and_register_then_offline_reuse(tmp_path, monkeypatch):
    root = _make_repo(tmp_path)
    calls: list[str] = []
    _preexisting = _manifest_artifact_ids(root)
    _install_net(monkeypatch, {KSH_GDP_URL: GDP_CSV.encode("utf-8")}, calls)

    report = ingest_source(root, IngestionRequest("ksh_gdp", "1960", "1985"))
    assert calls == [KSH_GDP_URL]
    assert report.downloaded and not report.cache_hit
    # one artifact per variable per use split (3 variables x 3 splits)
    assert len(report.artifact_ids) == 9
    assert {aid.split("__")[2] for aid in report.artifact_ids} == {"1960", "1970", "1979"}
    for path in report.artifact_paths:
        assert (root / path).is_file()
    # The repo manifest may already carry registered fixtures; this ingestion
    # must add exactly its own artifact ids and nothing else.
    assert _manifest_artifact_ids(root) == _preexisting | set(report.artifact_ids)

    # Offline re-run must hit the hash-verified cache, add nothing, and
    # never touch the network.
    report2 = ingest_source(root, IngestionRequest("ksh_gdp", "1960", "1985", offline=True))
    assert calls == [KSH_GDP_URL]
    assert report2.cache_hit and not report2.downloaded
    assert report2.artifact_ids == report.artifact_ids
    assert _manifest_artifact_ids(root) == _preexisting | set(report.artifact_ids)

    # The registered targets are loadable through the real registry gates.
    from hsm_lab.data.registry import DataRegistry

    registry = DataRegistry.load(root)
    artifact = next(
        a
        for a in registry.artifacts.values()
        if a.variable_id == "gdp" and "1960__1969" in a.artifact_id
    )
    frame = registry.load_targets(artifact, TargetUse.CALIBRATION)
    assert frame.height == 10
    assert frame["value"].to_list()[0] == 100.0


def test_refresh_with_changed_hash_creates_new_content_dir(tmp_path, monkeypatch):
    root = _make_repo(tmp_path)
    calls: list[str] = []
    _install_net(monkeypatch, {KSH_GDP_URL: GDP_CSV.encode("utf-8")}, calls)
    _preexisting = _manifest_artifact_ids(root)
    first = ingest_source(root, IngestionRequest("ksh_gdp", "1960", "1985"))

    # Same hash on refresh: bytes reused, manifest unchanged.
    second = ingest_source(root, IngestionRequest("ksh_gdp", "1960", "1985", refresh=True))
    assert second.downloaded and second.raw_sha256 == first.raw_sha256
    assert _manifest_artifact_ids(root) == _preexisting | set(first.artifact_ids)

    # Changed vintage: new content dir + hash-suffixed artifacts, old kept.
    revised = GDP_CSV.replace("1961;101;", "1961;110;")
    _install_net(monkeypatch, {KSH_GDP_URL: revised.encode("utf-8")}, calls)
    third = ingest_source(root, IngestionRequest("ksh_gdp", "1960", "1985", refresh=True))
    assert third.raw_sha256 != first.raw_sha256
    assert third.raw_path.parent != first.raw_path.parent
    assert first.raw_path.is_file()  # old vintage never overwritten
    assert _manifest_artifact_ids(root) == (
        _preexisting | set(first.artifact_ids) | set(third.artifact_ids)
    )
    assert {aid.split("__")[4] for aid in third.artifact_ids} == {third.raw_sha256[:12]}


def test_schema_drift_publishes_nothing_but_keeps_raw(tmp_path, monkeypatch):
    root = _make_repo(tmp_path)
    _preexisting = _manifest_artifact_ids(root)
    calls: list[str] = []
    drifted = GDP_CSV.replace("Gross domestic product (GDP)", "Gross domestic stuff")
    _install_net(monkeypatch, {KSH_GDP_URL: drifted.encode("utf-8")}, calls)
    with pytest.raises(SourceSchemaChanged, match="ksh_gdp"):
        ingest_source(root, IngestionRequest("ksh_gdp", "1960", "1985"))
    # Drift publishes nothing: the manifest keeps exactly its prior entries.
    assert _manifest_artifact_ids(root) == _preexisting
    # The drifted raw payload stays on disk for audit.
    content_dirs = list((root / "data" / "raw" / "ksh_gdp").iterdir())
    assert len(content_dirs) == 1
    assert (content_dirs[0] / "payload.csv").is_file()


def test_forcing_sources_register_single_window_artifact(tmp_path, monkeypatch):
    root = _make_repo(tmp_path)
    mnb_url = (
        "https://www.mnb.hu/en/arfolyam-tablazat?query&deviza=rbCurrencySelect"
        "&devizaSelected=USD&datefrom=01%2F01%2F1985&datetill=31%2F12%2F1985"
    )
    _install_net(monkeypatch, {mnb_url: mnb_html().encode("utf-8")}, [])
    report = ingest_source(root, IngestionRequest("mnb_official_fx", "1985-01-01", "1985-12-31"))
    assert report.variable_ids == ("official_fx",)
    assert len(report.artifact_ids) == 1
    from hsm_lab.data.registry import DataRegistry

    registry = DataRegistry.load(root)
    artifact = registry.artifacts[report.artifact_ids[0]]
    frame = registry.load_forcing(artifact)
    assert frame.height == 2
    assert frame["period"].to_list() == ["1985-01-03", "1985-03-30"]



# ----------------------------------------------------------------------
# Frequency / registry extensions used by the ingestion wave
# ----------------------------------------------------------------------


def test_dated_to_quarterly_takes_last_actual_quote():
    import polars as pl

    from hsm_lab.data.frequency import dated_to_quarterly

    frame = pl.DataFrame(
        {
            "period": ["1985-01-03", "1985-02-15", "1985-03-30", "1985-04-05"],
            "value": [51.30, 50.00, 50.88, 49.99],
        }
    )
    result = dated_to_quarterly(frame, "end_of_period").sort("period")
    assert result["period"].to_list() == ["1985Q1", "1985Q2"]
    assert result["value"].to_list() == [50.88, 49.99]  # last actual quote per quarter


def test_dated_to_quarterly_refuses_averaging_quotes():
    import polars as pl

    from hsm_lab.data.frequency import dated_to_quarterly

    frame = pl.DataFrame({"period": ["1985-01-03"], "value": [51.30]})
    with pytest.raises(ConfigError, match="end_of_period"):
        dated_to_quarterly(frame, "mean")


def test_cmea_transformation_coherence_gates():
    from hsm_lab.data.schemas import Transformation

    Transformation(
        method="cmea_five_year_contract",
        source_frequency="monthly",
        result_frequency="quarterly",
        aggregation="piecewise_constant",
        description="five-year contract rule",
    )
    Transformation(
        method="dated_to_quarterly",
        source_frequency="dated",
        result_frequency="quarterly",
        aggregation="end_of_period",
        description="last quote per quarter",
    )
    from pydantic import ValidationError

    with pytest.raises(ValidationError, match="piecewise_constant"):
        Transformation(
            method="cmea_five_year_contract",
            source_frequency="monthly",
            result_frequency="quarterly",
            aggregation="mean",
            description="bad",
        )
    with pytest.raises(ValidationError, match="monthly"):
        Transformation(
            method="cmea_five_year_contract",
            source_frequency="quarterly",
            result_frequency="quarterly",
            aggregation="piecewise_constant",
            description="bad",
        )


def test_cmea_method_is_reserved_for_cmea_variable(tmp_path):
    from tests.conftest import make_registry, source_spec, variable_spec, write_long_parquet

    rows = [
        {"variable_id": "world_oil_price", "period": p, "region_id": None,
         "commodity_id": None, "value": float(i + 1), "unit": "USD_per_barrel"}
        for i, p in enumerate(["1960Q1", "1960Q2", "1960Q3", "1960Q4"])
    ]
    digest = write_long_parquet(tmp_path / "fixtures" / "oil_q.parquet", rows)
    from tests.conftest import provenance

    prov = provenance(
        variable="world_oil_price",
        role="exogenous_forcing",
        evidence_status="reconstructed",
        original_frequency="monthly",
        model_frequency="quarterly",
        method="cmea_five_year_contract",
        aggregation="piecewise_constant",
        raw_sha256=digest,
        artifact_sha256=digest,
    )
    artifact = {
        "artifact_id": "oil_cmea",
        "variable_id": "world_oil_price",
        "path": "fixtures/oil_q.parquet",
        "source_id": "src_fixture",
        "role": "exogenous_forcing",
        "period": {"start": "1960Q1", "end": "1960Q4", "frequency": "quarterly"},
        "provenance": prov,
    }
    registry = make_registry(
        tmp_path,
        sources=[source_spec()],
        variables=[
            variable_spec(
                "world_oil_price",
                "external",
                "exogenous_forcing",
                unit="USD_per_barrel",
                original_frequency="monthly",
                aggregation="piecewise_constant",
            )
        ],
        artifacts=[artifact],
    )
    with pytest.raises(ResearchGateError, match="reserved.*cmea_energy_price"):
        registry.load_forcing(registry.artifacts["oil_cmea"])
